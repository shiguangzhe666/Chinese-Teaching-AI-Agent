"""
教案导出模块
支持导出为Word文档(.docx)和Markdown文件
"""
import io
import re
from datetime import datetime


def export_to_markdown(content, grade, lesson, plan_type="教案"):
    """导出为Markdown格式"""
    header = f"""# {grade}《{lesson}》{plan_type}

> 生成时间：{datetime.now().strftime('%Y年%m月%d日 %H:%M')}
> 语文AI教学智能体

---

"""
    return header + content


def markdown_to_docx(content, grade, lesson, plan_type="教案"):
    """
    将Markdown内容转换为Word文档
    使用python-docx库
    """
    try:
        from docx import Document
        from docx.shared import Pt, Cm, RGBColor
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml.ns import qn
    except ImportError:
        return None, "需要安装python-docx库：pip install python-docx"

    doc = Document()

    # 设置默认字体
    style = doc.styles['Normal']
    style.font.name = '宋体'
    style.font.size = Pt(11)
    style.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')

    # 标题
    title = doc.add_heading(f'{grade}《{lesson}》{plan_type}', level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # 副标题
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = meta.add_run(f'生成时间：{datetime.now().strftime("%Y年%m月%d日")}')
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(128, 128, 128)

    doc.add_paragraph()  # 空行

    # 解析Markdown内容
    lines = content.split('\n')
    for line in lines:
        line = line.rstrip()
        if not line:
            continue

        if line.startswith('## '):
            heading = doc.add_heading(line[3:], level=2)
            for run in heading.runs:
                run.font.name = '黑体'
                run.element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
        elif line.startswith('### '):
            heading = doc.add_heading(line[4:], level=3)
            for run in heading.runs:
                run.font.name = '黑体'
                run.element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
        elif line.startswith('#### '):
            heading = doc.add_heading(line[5:], level=4)
        elif line.startswith('- ') or line.startswith('* '):
            p = doc.add_paragraph(line[2:], style='List Bullet')
        elif re.match(r'^\d+\.', line):
            p = doc.add_paragraph(line, style='List Number')
        elif line.startswith('【') and line.endswith('】'):
            p = doc.add_paragraph()
            run = p.add_run(line)
            run.bold = True
            run.font.color.rgb = RGBColor(0, 100, 0)
        elif line.startswith('>'):
            p = doc.add_paragraph(line[1:].strip())
            p.paragraph_format.left_indent = Cm(1)
            for run in p.runs:
                run.font.color.rgb = RGBColor(100, 100, 100)
        elif line.startswith('---'):
            doc.add_paragraph('_' * 40)
        else:
            doc.add_paragraph(line)

    # 保存到内存
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer, None


def get_download_data(content, grade, lesson, fmt="markdown", plan_type="教案"):
    """获取下载数据"""
    if fmt == "docx":
        buffer, error = markdown_to_docx(content, grade, lesson, plan_type)
        if error:
            return None, error, None
        filename = f"{grade}_{lesson}_{plan_type}.docx"
        return buffer.getvalue(), None, filename
    else:
        md_content = export_to_markdown(content, grade, lesson, plan_type)
        filename = f"{grade}_{lesson}_{plan_type}.md"
        return md_content.encode('utf-8'), None, filename
